"""Gera o dashboard_academia.html com os dados da planilha embutidos.

Uso (na raiz do repositório):
    python ferramentas/build_painel.py

Lê as abas brutas de Modelo_Gestao_Academia_V2_Executivo.xlsx. A aba opcional
NotasFiscais vem da planilha quando existir; senão, das notas de exemplo
(ferramentas/notas_fiscais_exemplo.json).
"""
import datetime
import json
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
MODELO = RAIZ / "Modelo_Gestao_Academia_V2_Executivo.xlsx"
TEMPLATE = RAIZ / "ferramentas" / "template_painel.html"
NOTAS_EXEMPLO = RAIZ / "ferramentas" / "notas_fiscais_exemplo.json"
SAIDA = RAIZ / "dashboard_academia.html"
ABAS = ["Premissas", "Alunos", "Turmas", "Receitas", "Despesas", "Leads",
        "Churn_Cancelamentos", "Ocupacao_Horarios", "NotasFiscais"]


def valor(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def main():
    wb = openpyxl.load_workbook(MODELO, data_only=True)
    dados = {}
    for aba in ABAS:
        if aba not in wb.sheetnames:
            continue
        linhas = []
        for linha in wb[aba].iter_rows(values_only=True):
            linha = [valor(v) for v in linha]
            while linha and linha[-1] is None:
                linha.pop()
            if linha:
                linhas.append(linha)
        dados[aba] = linhas
    if "NotasFiscais" not in dados and NOTAS_EXEMPLO.exists():
        dados["NotasFiscais"] = json.loads(NOTAS_EXEMPLO.read_text(encoding="utf-8"))

    payload = json.dumps(dados, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = TEMPLATE.read_text(encoding="utf-8")
    assert "/*__EMBEDDED_DATA__*/null" in html, "marcador de dados não encontrado no template"
    SAIDA.write_text(html.replace("/*__EMBEDDED_DATA__*/null", payload), encoding="utf-8")
    print(f"ok: {SAIDA.name} ({SAIDA.stat().st_size // 1024} KB) · " + ", ".join(f"{k}: {len(v) - 1}" for k, v in dados.items()))


if __name__ == "__main__":
    main()
