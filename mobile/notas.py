"""Notas fiscais: armazenamento, leitura do XML da NF-e/NFS-e e exportação para Excel.

As notas lançadas no app ficam em `mobile/dados/notas_fiscais.csv`. Na primeira
execução o arquivo é criado com as notas de exemplo do painel.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import threading
import uuid
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

from calculos import mkey, norm, num, parse_data, parse_mes, txt

PASTA = Path(__file__).resolve().parent / "dados"
ARQUIVO = PASTA / "notas_fiscais.csv"
EXEMPLO = Path(__file__).resolve().parent.parent / "ferramentas" / "notas_fiscais_exemplo.json"
COLS = ["Data Emissão", "Número", "Série", "Tipo", "Emitente/Destinatário", "CNPJ/CPF", "Descrição", "Categoria",
        "Valor", "Status", "Chave de Acesso", "Mês"]
_trava = threading.Lock()


def normalizar(n: dict, origem: str) -> dict:
    d = n.get("data")
    d = d if isinstance(d, dt.date) else parse_data(d)
    chave = "".join(c for c in txt(n.get("chave")) if c.isdigit())
    return {"id": n.get("id") or uuid.uuid4().hex[:10], "origem": origem, "data": d,
            "mes": parse_mes(n.get("mes")) or mkey(d), "numero": txt(n.get("numero")), "serie": txt(n.get("serie")),
            "tipo": "Saída" if norm(n.get("tipo")).startswith("sai") else "Entrada",
            "parceiro": txt(n.get("parceiro")) or "—", "doc": "".join(c for c in txt(n.get("doc")) if c.isdigit()),
            "descricao": txt(n.get("descricao")), "categoria": txt(n.get("categoria")) or "Outros", "valor": num(n.get("valor")),
            "status": "Cancelada" if norm(n.get("status")).startswith("cancel") else "Válida", "chave": chave}


def chave_unica(n: dict) -> str:
    return "k" + n["chave"] if len(n["chave"]) == 44 else "|".join(["n", n["tipo"], n["numero"], n["serie"], norm(n["parceiro"])])


def _de_linha(l: dict) -> dict:
    return {"id": l.get("id"), "data": l.get("Data Emissão"), "numero": l.get("Número"), "serie": l.get("Série"),
            "tipo": l.get("Tipo"), "parceiro": l.get("Emitente/Destinatário"), "doc": l.get("CNPJ/CPF"),
            "descricao": l.get("Descrição"), "categoria": l.get("Categoria"), "valor": l.get("Valor"),
            "status": l.get("Status"), "chave": l.get("Chave de Acesso"), "mes": l.get("Mês")}


def _para_linha(n: dict) -> list:
    m = n["mes"]
    return [n["data"].strftime("%d/%m/%Y") if n["data"] else "", n["numero"], n["serie"], n["tipo"],
            "" if n["parceiro"] == "—" else n["parceiro"], fmt_doc(n["doc"]), n["descricao"], n["categoria"],
            round(n["valor"], 2), n["status"], n["chave"], f"{m[5:]}/{m[:4]}" if m else ""]


def _semear():
    PASTA.mkdir(parents=True, exist_ok=True)
    linhas = json.loads(EXEMPLO.read_text(encoding="utf-8"))[1:] if EXEMPLO.exists() else []
    with ARQUIVO.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id"] + COLS)
        for l in linhas:
            w.writerow([uuid.uuid4().hex[:10]] + l)


def carregar() -> list[dict]:
    with _trava:
        if not ARQUIVO.exists():
            _semear()
        with ARQUIVO.open(newline="", encoding="utf-8") as f:
            return [normalizar(_de_linha(l), "app") for l in csv.DictReader(f)]


def _gravar(notas: list[dict]):
    with ARQUIVO.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id"] + COLS)
        for n in notas:
            w.writerow([n["id"]] + _para_linha(n))


def adicionar(nova: dict) -> str | None:
    """Grava a nota; devolve mensagem de erro ou None."""
    atuais = carregar()
    if any(chave_unica(n) == chave_unica(nova) for n in atuais):
        return "Esta nota já está lançada (mesmo número, série e emitente ou mesma chave)."
    with _trava:
        _gravar(atuais + [nova])
    return None


def excluir(id_: str):
    atuais = carregar()
    with _trava:
        _gravar([n for n in atuais if n["id"] != id_])


def mesclar(*listas) -> list[dict]:
    vistos, out = set(), []
    for n in (x for l in listas for x in l):
        k = chave_unica(n)
        if k not in vistos:
            vistos.add(k)
            out.append(n)
    return out


def fmt_doc(d: str) -> str:
    if len(d) == 14:
        return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"
    if len(d) == 11:
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    return d


# ----------------------------------------------------------------------------
# XML da NF-e (modelos 55/65) e NFS-e padrão ABRASF / nacional
# ----------------------------------------------------------------------------
def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _achar(el, *caminho):
    for nome in caminho:
        if el is None:
            return None
        el = next((e for e in el.iter() if _local(e.tag) == nome and e is not el), None)
    return el


def _primeiro(*els):
    """Primeiro elemento existente (Element sem filhos é 'falso' num `or`, por isso a checagem explícita)."""
    return next((e for e in els if e is not None), None)


def _texto(el, *caminho) -> str:
    e = _achar(el, *caminho)
    return (e.text or "").strip() if e is not None else ""


def ler_xml(conteudo: bytes) -> dict:
    try:
        raiz = ET.fromstring(conteudo)
    except ET.ParseError as e:
        raise ValueError("O arquivo não é um XML válido.") from e
    inf = raiz if _local(raiz.tag) == "infNFe" else _achar(raiz, "infNFe")
    if inf is not None:
        emit = _achar(inf, "emit")
        prods = [(e.text or "").strip() for e in inf.iter() if _local(e.tag) == "xProd"]
        chave = _texto(raiz, "protNFe", "chNFe") or inf.get("Id", "")
        desc = prods[0] + (f" e mais {len(prods) - 1} ite{'ns' if len(prods) > 2 else 'm'}" if len(prods) > 1 else "") if prods else ""
        return {"data": (_texto(inf, "ide", "dhEmi") or _texto(inf, "ide", "dEmi"))[:10], "numero": _texto(inf, "ide", "nNF"),
                "serie": _texto(inf, "ide", "serie"), "parceiro": _texto(emit, "xNome"),
                "doc": _texto(emit, "CNPJ") or _texto(emit, "CPF"), "valor": _texto(inf, "total", "ICMSTot", "vNF"),
                "descricao": desc, "chave": "".join(c for c in chave if c.isdigit()), "tipo": "Entrada"}
    nfse = _primeiro(_achar(raiz, "InfNfse"), _achar(raiz, "infNFSe"))
    if nfse is not None:
        prest = _primeiro(_achar(nfse, "PrestadorServico"), _achar(nfse, "Prestador"), _achar(nfse, "emit"))
        return {"data": (_texto(nfse, "DataEmissao") or _texto(nfse, "dhEmi") or _texto(nfse, "dhProc"))[:10],
                "numero": _texto(nfse, "Numero") or _texto(nfse, "nNFSe"), "serie": "NFS-e",
                "parceiro": _texto(prest, "RazaoSocial") or _texto(prest, "xNome"),
                "doc": _texto(prest, "Cnpj") or _texto(prest, "CNPJ") or _texto(prest, "Cpf"),
                "valor": _texto(nfse, "ValorServicos") or _texto(nfse, "vServ") or _texto(nfse, "vLiq"),
                "descricao": _texto(nfse, "Discriminacao") or _texto(nfse, "xDescServ"), "chave": "", "tipo": "Entrada"}
    raise ValueError("Não reconheci o layout deste XML. Preencha os campos manualmente.")


# ----------------------------------------------------------------------------
# Exportação: planilha completa (abas brutas + NotasFiscais), compatível com o painel
# ----------------------------------------------------------------------------
def exportar_xlsx(book: dict[str, list[list]], notas: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, linhas in book.items():
        if norm(nome).replace(" ", "").replace("_", "") == "notasfiscais":
            continue
        ws = wb.create_sheet(nome[:31])
        for l in linhas:
            ws.append([v.strftime("%d/%m/%Y") if isinstance(v, (dt.date, dt.datetime)) else v for v in l])
    ws = wb.create_sheet("NotasFiscais")
    ws.append(COLS)
    for n in sorted(notas, key=lambda n: n["data"] or dt.date.min):
        ws.append(_para_linha(n))
    for col, larg in zip("ABCDEFGHIJKL", [12, 10, 7, 9, 32, 20, 36, 14, 12, 10, 46, 9]):
        ws.column_dimensions[col].width = larg
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
